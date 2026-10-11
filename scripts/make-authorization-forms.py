# Run: python3 scripts/make-authorization-forms.py public/forms  (needs reportlab, python-docx, DejaVu fonts)
# Letter of authorization forms (EN/HU) for RateThatJob access requests: fillable PDF + DOCX.
import sys
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor, white
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

OUT = sys.argv[1]
pdfmetrics.registerFont(TTFont("DV", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"))
pdfmetrics.registerFont(TTFont("DVB", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"))
INK = HexColor("#14161A"); SLATE = HexColor("#4A4F57"); LINE = HexColor("#B5BAC1")

T = {
 "en": {
  "title": "Letter of authorization",
  "subtitle": "RateThatJob employer profile (ratethatjob.com)",
  "intro": "The company below authorizes the person named in section 2 to manage its RateThatJob employer profile on its behalf for the workplace(s) listed.",
  "s1": "1. Company",
  "company": "Company name", "regno": "Company registration / trade license number", "office": "Registered office (address)",
  "locations": "Workplace(s) concerned (name and address of each location)",
  "s2": "2. Authorized person",
  "person": "Full name", "email": "Email address (used to sign in to RateThatJob)", "position": "Position at the company",
  "s3": "3. Scope of the authorization",
  "scope": ["Manage the company's RateThatJob profile for the workplace(s) above",
            "Reply to reviews and dispute unverified reviews",
            "Invite staff to leave reviews",
            "Post and manage job ads and handle applications",
            "Order and manage Pro subscriptions and job ad packages, and pay for them"],
  "s4": "4. Validity",
  "from": "Valid from", "until": "Valid until (leave empty: until revoked in writing)",
  "s5": "5. Signed for the company by its legal representative",
  "rep": "Name of the legal representative", "repPosition": "Position", "place": "Place", "date": "Date",
  "signature": "Signature", "stamp": "Company stamp (if any)",
  "footer": "Upload the signed form with a photo ID of the authorized person when you request access at ratethatjob.com/business/claim. RateThatJob uses it only to check the request and deletes it after the decision.",
 },
 "hu": {
  "title": "Meghatalmazás",
  "subtitle": "RateThatJob munkáltatói profil (ratethatjob.com)",
  "intro": "Az alább megnevezett cég meghatalmazza a 2. pontban megjelölt személyt, hogy a nevében kezelje a cég RateThatJob munkáltatói profilját a felsorolt munkahely(ek)re vonatkozóan.",
  "s1": "1. A cég adatai",
  "company": "Cégnév", "regno": "Cégjegyzékszám / trade license szám", "office": "Székhely (cím)",
  "locations": "Érintett munkahely(ek) (minden telephely neve és címe)",
  "s2": "2. A meghatalmazott",
  "person": "Teljes név", "email": "E-mail-cím (amellyel a RateThatJobra belép)", "position": "Beosztás a cégnél",
  "s3": "3. A meghatalmazás tárgya",
  "scope": ["A cég RateThatJob profiljának kezelése a fenti munkahely(ek)re",
            "Válaszadás az értékelésekre, az igazolatlan értékelések vitatása",
            "Munkatársak meghívása értékelésre",
            "Álláshirdetések feladása, kezelése és a jelentkezések kezelése",
            "Pro előfizetés és álláshirdetési csomagok megrendelése, kezelése és kifizetése"],
  "s4": "4. Érvényesség",
  "from": "Érvényes ettől", "until": "Érvényes eddig (üresen hagyva: írásbeli visszavonásig)",
  "s5": "5. A cég nevében aláírja a törvényes képviselő",
  "rep": "A törvényes képviselő neve", "repPosition": "Beosztás", "place": "Kelt (hely)", "date": "Dátum",
  "signature": "Aláírás", "stamp": "Cégbélyegző (ha van)",
  "footer": "Az aláírt nyomtatványt a meghatalmazott személyazonosító okmányával együtt töltsd fel a hozzáférés igénylésekor: ratethatjob.com/business/claim. A RateThatJob csak az igénylés ellenőrzésére használja, és a döntés után törli.",
 },
}

def wrap(c, text, font, size, width):
    words, lines, line = text.split(), [], ""
    for w in words:
        test = f"{line} {w}".strip()
        if c.stringWidth(test, font, size) > width and line:
            lines.append(line); line = w
        else:
            line = test
    if line: lines.append(line)
    return lines

def pdf(lang):
    t = T[lang]
    path = f"{OUT}/rtj-authorization-{lang}.pdf"
    c = canvas.Canvas(path, pagesize=A4)
    c.setTitle(f"{t['title']} – RateThatJob"); c.setAuthor("RateThatJob")
    W, H = A4; L = 50; R = W - 50; y = H - 60
    form = c.acroForm
    c.setFillColor(INK); c.setFont("DVB", 20); c.drawString(L, y, t["title"]); y -= 20
    c.setFont("DV", 11); c.setFillColor(SLATE); c.drawString(L, y, t["subtitle"]); y -= 22
    c.setFillColor(INK); c.setFont("DV", 10)
    for ln in wrap(c, t["intro"], "DV", 10, R - L): c.drawString(L, y, ln); y -= 14
    y -= 6
    n = [0]
    def section(title):
        nonlocal y
        y -= 8; c.setFont("DVB", 12); c.setFillColor(INK); c.drawString(L, y, title); y -= 6
        c.setStrokeColor(LINE); c.line(L, y, R, y); y -= 14
    def field(label, x=L, w=None, h=20, multiline=False, advance=True):
        nonlocal y
        w = w or (R - L)
        c.setFont("DV", 8.5); c.setFillColor(SLATE); c.drawString(x, y, label)
        n[0] += 1
        form.textfield(name=f"f{n[0]}", tooltip=label, x=x, y=y - 4 - h, width=w, height=h,
                       borderColor=LINE, fillColor=white, textColor=INK, borderWidth=0.8, fontName="Helvetica",
                       fontSize=0 if multiline else 10, fieldFlags="multiline" if multiline else "")
        if advance: y -= h + 18
    half = (R - L - 12) / 2
    section(t["s1"])
    field(t["company"]); field(t["regno"]); field(t["office"]); field(t["locations"], h=46, multiline=True)
    section(t["s2"])
    field(t["person"]); field(t["email"], w=half, advance=False); field(t["position"], x=L + half + 12, w=half)
    section(t["s3"])
    c.setFont("DV", 9.5); c.setFillColor(INK)
    for item in t["scope"]:
        c.drawString(L + 4, y, "•"); 
        for i, ln in enumerate(wrap(c, item, "DV", 9.5, R - L - 18)): c.drawString(L + 16, y, ln); y -= 13
    y -= 4
    section(t["s4"])
    field(t["from"], w=half, advance=False); field(t["until"], x=L + half + 12, w=half)
    section(t["s5"])
    field(t["rep"], w=half, advance=False); field(t["repPosition"], x=L + half + 12, w=half)
    field(t["place"], w=half, advance=False); field(t["date"], x=L + half + 12, w=half)
    # Signature and stamp: signed by hand.
    c.setFont("DV", 8.5); c.setFillColor(SLATE)
    c.drawString(L, y, t["signature"]); c.drawString(L + half + 12, y, t["stamp"])
    c.setStrokeColor(LINE); c.rect(L, y - 74, half, 66); c.rect(L + half + 12, y - 74, half, 66)
    y -= 92
    c.setFont("DV", 8); c.setFillColor(SLATE)
    for ln in wrap(c, t["footer"], "DV", 8, R - L): c.drawString(L, y, ln); y -= 11
    c.save()

def docx(lang):
    t = T[lang]
    d = Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Cm(2); s.top_margin = s.bottom_margin = Cm(1.4)
    st = d.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(9.5)
    h = d.add_paragraph(); r = h.add_run(t["title"]); r.bold = True; r.font.size = Pt(20)
    p = d.add_paragraph(); r = p.add_run(t["subtitle"]); r.font.color.rgb = RGBColor(0x4A, 0x4F, 0x57)
    d.add_paragraph(t["intro"])
    def section(title):
        p = d.add_paragraph(); p.paragraph_format.space_before = Pt(6); r = p.add_run(title); r.bold = True; r.font.size = Pt(12)
    def table(rows, tall=()):
        tb = d.add_table(rows=0, cols=2); tb.style = "Table Grid"; tb.alignment = WD_TABLE_ALIGNMENT.CENTER
        for label in rows:
            row = tb.add_row()
            row.height = Cm(1.5 if label in tall else 0.68)
            row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            cells = row.cells
            cells[0].width = Cm(6.5); cells[1].width = Cm(10.5)
            cells[0].text = label
            cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        return tb
    section(t["s1"]); table([t["company"], t["regno"], t["office"], t["locations"]], tall=(t["locations"],))
    section(t["s2"]); table([t["person"], t["email"], t["position"]])
    section(t["s3"])
    for item in t["scope"]: d.add_paragraph(item, style="List Bullet")
    section(t["s4"]); table([t["from"], t["until"]])
    section(t["s5"]); table([t["rep"], t["repPosition"], t["place"], t["date"], t["signature"], t["stamp"]], tall=(t["signature"], t["stamp"]))
    p = d.add_paragraph(); p.paragraph_format.space_before = Pt(12)
    r = p.add_run(t["footer"]); r.font.size = Pt(8); r.font.color.rgb = RGBColor(0x4A, 0x4F, 0x57)
    d.core_properties.title = f"{t['title']} – RateThatJob"; d.core_properties.author = "RateThatJob"
    d.save(f"{OUT}/rtj-authorization-{lang}.docx")

for lang in ("en", "hu"):
    pdf(lang); docx(lang)
print("ok")
